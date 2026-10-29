
import re
from pathlib import Path
from typing import Union

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

# Common headings and terms found in resumes.
RESUME_SECTIONS = {
    "education": [
        "education",
        "academic background",
        "educational qualifications",
        "academic qualifications",
    ],
    "skills": [
        "skills",
        "technical skills",
        "core competencies",
        "technologies",
        "technical expertise",
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "internship",
        "internships",
    ],
    "projects": [
        "projects",
        "academic projects",
        "personal projects",
        "project experience",
    ],
}

# Phrases that commonly indicate an academic or project report.
REPORT_INDICATORS = [
    "table of contents",
    "certificate of completion",
    "declaration",
    "literature review",
    "research methodology",
    "abstract",
    "chapter 1",
    "chapter 2",
    "problem statement",
    "proposed methodology",
    "experimental results",
    "conclusion and future scope",
]


def extract_text(file_path: Union[str, Path]) -> str:
    """
    Extract text from a PDF or DOCX file.

    Args:
        file_path: Path to the uploaded document.

    Returns:
        Extracted document text.

    Raises:
        ValueError: If the file type is unsupported or no text is found.
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file format. Please upload a PDF or DOCX resume."
        )

    try:
        if extension == ".pdf":
            reader = PdfReader(str(path))

            if reader.is_encrypted:
                raise ValueError(
                    "This PDF is password-protected. Please upload an unlocked resume."
                )

            pages = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")

            text = "\n".join(pages)

        else:
            document = Document(str(path))
            paragraphs = [p.text for p in document.paragraphs]

            # Also extract text from tables, which resumes often use.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(paragraphs)

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Unable to read this document. Please upload a valid PDF or DOCX file."
        ) from exc

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if len(text) < 80:
        raise ValueError(
            "Very little text could be extracted. If this is a scanned PDF, "
            "please upload a text-based PDF or a DOCX resume."
        )

    return text


def validate_resume_text(text: str) -> tuple[bool, str]:
    """
    Check whether extracted document text has enough resume-like structure.

    This is a heuristic, not a perfect resume classifier.
    """
    normalized = re.sub(r"\s+", " ", text.lower()).strip()

    if len(normalized) < 80:
        return False, (
            "The document contains too little readable text. "
            "Please upload your resume in PDF or DOCX format."
        )

    # Count distinct resume section categories found in the document.
    matched_categories = 0

    for phrases in RESUME_SECTIONS.values():
        if any(
            re.search(
                r"(?<![a-z])" + re.escape(phrase) + r"(?![a-z])",
                normalized,
            )
            for phrase in phrases
        ):
            matched_categories += 1

    # Look for contact details and common resume-related information.
    has_email = bool(
        re.search(
            r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
            text,
            re.IGNORECASE,
        )
    )

    has_phone = bool(
        re.search(
            r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)",
            text,
        )
    )

    has_resume_terms = bool(
        re.search(
            r"\b(resume|curriculum vitae|linkedin|github|"
            r"bachelor|b\.tech|b\.e\.|computer science|"
            r"developer|engineer|programming|certifications)\b",
            normalized,
        )
    )

    report_indicator_count = sum(
        1 for phrase in REPORT_INDICATORS if phrase in normalized
    )

    # Reject documents that look strongly like academic reports and
    # have very little recognizable resume structure.
    if report_indicator_count >= 3 and matched_categories < 2:
        return False, (
            "This document appears to be an academic or project report, "
            "not a resume. Please upload your resume in PDF or DOCX format."
        )

    # Require multiple signals to reduce accidental acceptance of reports.
    resume_signals = sum(
        [
            matched_categories >= 2,
            has_email,
            has_phone,
            has_resume_terms,
        ]
    )

    if resume_signals < 2:
        return False, (
            "This document doesn't appear to contain a recognizable resume "
            "structure. Please upload your resume in PDF or DOCX format, "
            "including sections such as Skills, Education, Experience, "
            "or Projects."
        )

    return True, "Resume validation successful."


def parse_resume(file_path: Union[str, Path]) -> str:
    """
    Extract and validate a resume before it is sent for analysis.
    """
    text = extract_text(file_path)

    is_resume, message = validate_resume_text(text)

    if not is_resume:
        raise ValueError(message)

    return text


# Backward-compatible aliases for projects that use these function names.
def extract_resume_text(file_path: Union[str, Path]) -> str:
    return parse_resume(file_path)


def parse_document(file_path: Union[str, Path]) -> str:
    return parse_resume(file_path)
